"""ExportService infrastructure tests (Brief §6)."""

from __future__ import annotations

import io
import json

import pytest


ROWS = [
    {"business_name": "Alpha Traders", "city": "Mumbai", "phone": "+919820000001"},
    {"business_name": "Beta Industries", "city": "Delhi", "phone": "+919820000002"},
]


async def _export(app, fmt: str, rows=ROWS):
    from app.services.export import ExportService

    files = app.state.files
    export_service = ExportService(files)
    async with app.state.db.session() as session:
        record = await export_service.export(
            session,
            format_name=fmt,
            rows=rows,
            base_name="test-leads",
            metadata={"source": "unit-test"},
        )
    return record


async def test_csv_export_roundtrip(app):
    record = await _export(app, "csv")
    assert record.category == "EXPORT"
    assert record.metadata_json["format"] == "csv"
    assert record.metadata_json["row_count"] == 2

    import csv as csv_mod

    with app.state.storage.open(record.path, category="EXPORT") as fh:
        text = fh.read().decode("utf-8")
    parsed = list(csv_mod.DictReader(io.StringIO(text)))
    assert parsed[0]["business_name"] == "Alpha Traders"
    assert parsed[1]["city"] == "Delhi"


async def test_json_export_roundtrip(app):
    record = await _export(app, "json")
    with app.state.storage.open(record.path, category="EXPORT") as fh:
        data = json.loads(fh.read().decode("utf-8"))
    assert len(data) == 2
    assert data[1]["business_name"] == "Beta Industries"


async def test_xlsx_export_roundtrip(app):
    from openpyxl import load_workbook

    record = await _export(app, "xlsx")
    with app.state.storage.open(record.path, category="EXPORT") as fh:
        content = fh.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb["export"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("business_name", "city", "phone")
    assert rows[1][0] == "Alpha Traders"


async def test_unsupported_format_rejected(app):
    from app.core.errors import NotFoundError

    with pytest.raises(NotFoundError):
        await _export(app, "pdf")


async def test_export_registers_file_metadata(app):
    record = await _export(app, "csv")
    async with app.state.db.session() as session:
        from sqlalchemy import select

        from app.models.file import FileRecord

        row = await session.scalar(select(FileRecord).where(FileRecord.id == record.id))
    assert row is not None
    assert row.size > 0
    assert row.checksum_sha256
