"""Lead import system (Phase 4 §20-§25): CSV / XLSX / JSON / JSONL.

Flow: Upload → Inspect → Map columns → Preview → Validate → Import
      → Deduplicate → Complete (+ downloadable rejected-rows report)

Design:
- rows are STREAMED (csv reader / openpyxl read-only / jsonl lines); only
  whole-file JSON needs a full parse (inherent to the format, documented)
- invalid rows never disappear silently: counters + error summary + a
  rejected-rows CSV stored via StorageService
- duplicate strategies: SKIP_DUPLICATES (default) | UPDATE_EXISTING |
  CREATE_NEW | REVIEW (creates review candidates, never auto-merges)
- import runs in chunks (commit per chunk); large batches are claimed by the
  background worker (DB-first claim, same pattern as the scrape queue)
"""

from __future__ import annotations

import csv
import io
import json
import os
import tempfile
import uuid
from datetime import datetime, timezone
from typing import AsyncIterator, BinaryIO, Iterator

from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import NotFoundError, ValidationError
from app.core.logging import get_logger
from app.models.lead import DuplicateConfidence, ImportBatch, ImportStatus, LeadStatus
from app.models.scrape import Lead
from app.services.files import FileService
from app.services.leads.activity import EVENT_IMPORTED, LeadActivityService
from app.services.leads.dedup_engine import DuplicateDetectionService
from app.services.leads.normalization import normalize_lead_payload
from app.services.leads.quality import compute_quality_score
from app.services.leads.tags import TagService
from app.services.scraping.dedup import Deduplicator, MatchConfidence
from app.services.storage import StorageService

logger = get_logger("qbit.leads.import")

#: columns a user may map into (whitelist)
MAPPABLE_FIELDS = (
    "business_name", "contact_name", "first_name", "last_name", "email", "phone",
    "website", "address", "city", "state", "postal_code", "country", "category",
    "industry", "source_id", "source_url", "source",
)

SPECIAL_FIELDS = ("tags",)  # mapped column split on ; or |

DUP_STRATEGIES = ("SKIP_DUPLICATES", "UPDATE_EXISTING", "CREATE_NEW", "REVIEW")
DUPLICATE_CONFIDENCES = {MatchConfidence.HIGH, MatchConfidence.MEDIUM}

ERROR_REPORT_LIMIT = 10000  # rejected rows persisted to the CSV report


def _review_confidence(match) -> str:
    """Map the pipeline MatchConfidence ladder onto the workspace ladder."""
    if match.confidence is MatchConfidence.MEDIUM:
        return DuplicateConfidence.MEDIUM.value
    if (match.matched_on or "") in ("email", "phone"):
        return DuplicateConfidence.EXACT.value
    return DuplicateConfidence.HIGH.value


class LeadImportService:
    def __init__(self, storage: StorageService, files: FileService) -> None:
        self.storage = storage
        self.files = files
        self.activities = LeadActivityService()
        self.tags = TagService()
        self.duplicates = DuplicateDetectionService()

    # ------------------------------------------------------------- lifecycle
    async def create_batch(
        self,
        session: AsyncSession,
        *,
        file_record,
        format_name: str,
        filename: str,
        created_by: uuid.UUID | None,
        options: dict | None = None,
    ) -> ImportBatch:
        format_name = (format_name or "").lower()
        if format_name not in ("csv", "xlsx", "json", "jsonl"):
            raise ValidationError(f"Unsupported import format: {format_name}")
        batch = ImportBatch(
            filename=(filename or file_record.name)[:260],
            file_id=file_record.id,
            format=format_name,
            status=ImportStatus.QUEUED.value,
            options=options or {},
            created_by=created_by,
        )
        session.add(batch)
        await session.commit()
        await session.refresh(batch)
        logger.info(
            "Import batch created",
            extra={"extra_fields": {"batch_id": str(batch.id), "format": format_name}},
        )
        return batch

    async def inspect(self, session: AsyncSession, batch: ImportBatch) -> dict:
        """Headers + sample rows + row estimate — feeds the mapping UI."""
        file_record = await self.files.get(session, batch.file_id) if batch.file_id else None
        if file_record is None:
            raise NotFoundError("Import file not found")
        stream = self.storage.open(file_record.path, category=file_record.category)
        with stream:
            if batch.format == "csv":
                return self._inspect_csv(stream, batch.options)
            if batch.format == "xlsx":
                return self._inspect_xlsx(stream, batch.options)
            if batch.format == "json":
                return self._inspect_json(stream)
            return self._inspect_jsonl(stream)

    # ------------------------------------------------------------ row feeds
    def _row_iterator(self, stream: BinaryIO, batch: ImportBatch) -> Iterator[dict]:
        if batch.format == "csv":
            yield from self._iter_csv(stream, batch.options)
        elif batch.format == "xlsx":
            yield from self._iter_xlsx(stream, batch.options)
        elif batch.format == "json":
            yield from self._iter_json(stream)
        else:
            yield from self._iter_jsonl(stream)

    def _sheet_of(self, workbook, options: dict):
        sheet = options.get("sheet")
        if sheet is None:
            return workbook[workbook.sheetnames[0]]
        if isinstance(sheet, int) or (isinstance(sheet, str) and sheet.isdigit()):
            index = int(sheet) - 1
            if index < 0 or index >= len(workbook.sheetnames):
                raise ValidationError(f"Sheet index out of range: {sheet}")
            return workbook[workbook.sheetnames[index]]
        if sheet in workbook.sheetnames:
            return workbook[sheet]
        raise ValidationError(f"Unknown sheet: {sheet}")

    @staticmethod
    def _cell(value) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _iter_xlsx(self, stream: BinaryIO, options: dict) -> Iterator[dict]:
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(stream, read_only=True, data_only=True)
        except Exception as exc:  # malformed workbooks
            raise ValidationError(f"Malformed XLSX: {exc}") from exc
        header_row = int(options.get("header_row", 1) or 1)
        sheet = self._sheet_of(workbook, options)
        headers: list[str] | None = None
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [self._cell(c) for c in row]
            if row_index < header_row:
                continue
            if row_index == header_row:
                headers = [c for c in cells if c] or [
                    f"column_{i + 1}" for i in range(len(cells))
                ]
                continue
            if all(not c for c in cells):
                continue
            yield dict(zip(headers, cells))
        try:
            workbook.close()
        except Exception:  # noqa: BLE001 — close is best-effort
            pass

    def _iter_csv(self, stream: BinaryIO, options: dict) -> Iterator[dict]:
        delimiter = (options.get("delimiter") or ",").strip() or ","
        text = io.TextIOWrapper(stream, encoding="utf-8-sig", errors="replace", newline="")
        reader = csv.DictReader(text, delimiter=delimiter)
        if reader.fieldnames is None:
            return
        headers = [(name or f"column_{i + 1}").strip() for i, name in enumerate(reader.fieldnames)]
        reader.fieldnames = headers
        yield from reader

    def _iter_json(self, stream: BinaryIO) -> Iterator[dict]:
        try:
            payload = json.loads(stream.read().decode("utf-8-sig"))
        except (ValueError, UnicodeDecodeError) as exc:
            raise ValidationError(f"Malformed JSON: {exc}") from exc
        if isinstance(payload, dict):
            payload = [payload]
        if not isinstance(payload, list):
            raise ValidationError("JSON import must be an array of objects")
        for item in payload:
            if isinstance(item, dict):
                yield {str(k): v for k, v in item.items()}

    def _iter_jsonl(self, stream: BinaryIO) -> Iterator[dict]:
        for line_number, raw in enumerate(stream, start=1):
            line = raw.decode("utf-8-sig", errors="replace").strip()
            if not line:
                continue
            try:
                item = json.loads(line)
            except ValueError as exc:
                raise ValidationError(f"Malformed JSON on line {line_number}: {exc}") from exc
            if isinstance(item, dict):
                yield {str(k): v for k, v in item.items()}

    # -------------------------------------------------------------- inspect
    def _inspect_csv(self, stream: BinaryIO, options: dict) -> dict:
        preview = self._iter_csv(stream, options)
        headers: list[str] = []
        sample = []
        total = 0
        for i, row in enumerate(preview):
            if i == 0:
                headers = list(row.keys())
            if len(sample) < 5:
                sample.append({k: v for k, v in row.items() if k in (headers or [])})
            total += 1
        return {"format": "csv", "columns": headers, "sample": sample, "total_rows": total}

    def _inspect_xlsx(self, stream: BinaryIO, options: dict) -> dict:
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(stream, read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError(f"Malformed XLSX: {exc}") from exc
        sheet = self._sheet_of(workbook, options)
        header_row = int(options.get("header_row", 1) or 1)
        headers: list[str] = []
        sample: list[dict] = []
        total = 0
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [self._cell(c) for c in row]
            if row_index == header_row:
                headers = [c for c in cells if c] or [f"column_{i + 1}" for i in range(len(cells))]
                continue
            if row_index < header_row or all(not c for c in cells):
                continue
            if len(sample) < 5:
                sample.append(dict(zip(headers, cells)))
            total += 1
        try:
            workbook.close()
        except Exception:  # noqa: BLE001
            pass
        return {
            "format": "xlsx",
            "sheets": workbook.sheetnames,
            "columns": headers,
            "sample": sample,
            "total_rows": total,
        }

    def _collect_json_columns(self, rows: list[dict]) -> list[str]:
        columns: list[str] = []
        for row in rows:
            for key in row:
                if key not in columns:
                    columns.append(key)
        return columns

    def _inspect_json(self, stream: BinaryIO) -> dict:
        rows = list(self._iter_json(stream))[:50]
        return {
            "format": "json",
            "columns": self._collect_json_columns(rows),
            "sample": rows[:5],
            "total_rows": len(rows),
        }

    def _inspect_jsonl(self, stream: BinaryIO) -> dict:
        rows = list(self._iter_jsonl(stream))[:50]
        return {
            "format": "jsonl",
            "columns": self._collect_json_columns(rows),
            "sample": rows[:5],
            "total_rows": len(rows),
        }

    # -------------------------------------------------------------- mapping
    @staticmethod
    def _validate_mapping(mapping: dict) -> dict[str, str]:
        clean: dict[str, str] = {}
        for column, field in (mapping or {}).items():
            field = str(field).strip()
            if not field:
                continue
            if field not in MAPPABLE_FIELDS and field not in SPECIAL_FIELDS:
                raise ValidationError(f"Cannot map into unknown lead field: {field}")
            clean[str(column)] = field
        if not clean:
            raise ValidationError("Column mapping is empty — map at least one column")
        return clean

    @staticmethod
    def _apply_mapping(row: dict, mapping: dict) -> tuple[dict, dict]:
        payload: dict = {}
        extra: dict = {}
        mapped_columns = set(mapping)
        for column, value in row.items():
            field = mapping.get(column)
            if field is None:
                if value not in (None, ""):
                    extra[str(column)] = value
                continue
            if field == "tags":
                tags = [t.strip() for t in str(value or "").replace("|", ";").split(";") if t.strip()]
                if tags:
                    payload["tags"] = tags
                continue
            payload[field] = value
        if extra:
            payload.setdefault("metadata", {})["extra_fields"] = extra
        return payload, extra

    # ------------------------------------------------------------ execution
    async def run(self, session: AsyncSession, batch: ImportBatch, *, worker_owner: str | None = None) -> ImportBatch:
        """Process a QUEUED batch end-to-end (also used by the worker)."""
        if batch.status not in (ImportStatus.QUEUED.value, ImportStatus.PROCESSING.value):
            raise ValidationError(f"Import batch is not runnable (status={batch.status})")

        now = datetime.now(timezone.utc)
        batch.status = ImportStatus.PROCESSING.value
        batch.started_at = batch.started_at or now
        if worker_owner:
            batch.lease_owner = worker_owner
            batch.leased_at = now
        await session.commit()

        mapping = self._validate_mapping(batch.mapping)
        options = batch.options or {}
        strategy = (options.get("duplicate_strategy") or "SKIP_DUPLICATES").upper()
        if strategy not in DUP_STRATEGIES:
            raise ValidationError(f"Unknown duplicate strategy: {strategy}")
        default_status = options.get("default_status") or LeadStatus.NEW.value
        import_tags = [str(t) for t in (options.get("tags") or []) if str(t).strip()]

        file_record = await self.files.get(session, batch.file_id)
        source_label = options.get("source_name") or f"import:{batch.filename}"

        rejected_path = None
        rejected_writer = None
        rejected_handle = None
        import tempfile

        dedup = Deduplicator()
        chunk: list[tuple[dict, dict]] = []  # (payload, raw_row)
        row_number = 0
        counters = {"total": 0, "valid": 0, "invalid": 0, "imported": 0,
                    "duplicate": 0, "updated": 0, "review": 0}
        error_summary: list[dict] = []

        try:
            rejected_handle = tempfile.NamedTemporaryFile(
                mode="w+", suffix=".csv", prefix=f"import-errors-{str(batch.id)[:8]}-",
                encoding="utf-8", delete=False,
            )
            rejected_writer = csv.writer(rejected_handle)
            rejected_writer.writerow(["row", "error", "data"])
            rejected_path = rejected_handle.name

            stream = self.storage.open(file_record.path, category=file_record.category)
            with stream:
                for raw_row in self._row_iterator(stream, batch):
                    row_number += 1
                    counters["total"] += 1
                    payload, _extra = self._apply_mapping(raw_row, mapping)
                    error = self._row_error(payload)
                    if error:
                        counters["invalid"] += 1
                        if len(error_summary) < 1000:
                            error_summary.append({"row": row_number, "error": error})
                        if counters["invalid"] <= ERROR_REPORT_LIMIT:
                            rejected_writer.writerow([row_number, error, json.dumps(raw_row, default=str, ensure_ascii=False)])
                        continue

                    clean, field_errors = normalize_lead_payload(payload)
                    if field_errors:
                        counters["invalid"] += 1
                        first_error = "; ".join(f"{k}: {v}" for k, v in field_errors.items())
                        if len(error_summary) < 1000:
                            error_summary.append({"row": row_number, "error": first_error})
                        if counters["invalid"] <= ERROR_REPORT_LIMIT:
                            rejected_writer.writerow([row_number, first_error, json.dumps(raw_row, default=str, ensure_ascii=False)])
                        continue

                    counters["valid"] += 1
                    tags = list(clean.pop("tags", []) or [])
                    tags.extend(import_tags)
                    outcome = await self._import_row(
                        session, batch, clean, tags, strategy, dedup,
                        default_status, source_label, file_record.id, counters,
                    )
                    _ = outcome

                    if counters["total"] % 500 == 0:
                        await self._flush(session, batch, counters, error_summary)
                await self._flush(session, batch, counters, error_summary)
        except ValidationError as exc:
            batch.status = ImportStatus.FAILED.value
            batch.error_summary = [{"row": row_number + 1, "error": exc.message}]
            batch.completed_at = datetime.now(timezone.utc)
            await session.commit()
            return batch
        finally:
            if rejected_handle is not None:
                rejected_handle.close()

        # --- rejected rows report -------------------------------------------
        if rejected_path is not None and counters["invalid"] > 0:
            try:
                with open(rejected_path, "rb") as handle:
                    report = await self.files.store(
                        session,
                        content=handle,
                        filename=f"import-errors-{str(batch.id)[:8]}.csv",
                        mime_type="text/csv",
                        category="IMPORT",
                        created_by=batch.created_by,
                        metadata={"import_batch_id": str(batch.id), "kind": "rejected_rows"},
                    )
                batch.error_file_id = report.id
            finally:
                try:
                    os.unlink(rejected_path)
                except OSError:  # pragma: no cover
                    pass

        batch.imported_rows = counters["imported"]
        batch.duplicate_rows = counters["duplicate"]
        batch.updated_rows = counters["updated"]
        batch.review_rows = counters["review"]
        batch.valid_rows = counters["valid"]
        batch.invalid_rows = counters["invalid"]
        batch.total_rows = counters["total"]
        batch.error_summary = error_summary[:1000]
        batch.completed_at = datetime.now(timezone.utc)
        batch.status = (
            ImportStatus.COMPLETED.value
            if counters["invalid"] == 0
            else ImportStatus.PARTIAL.value
        )
        await session.commit()
        logger.info(
            "Import batch finished",
            extra={"extra_fields": {
                "batch_id": str(batch.id), "status": batch.status,
                "imported": batch.imported_rows, "invalid": batch.invalid_rows,
                "duplicates": batch.duplicate_rows,
            }},
        )
        return batch

    async def _import_row(
        self, session: AsyncSession, batch: ImportBatch, clean: dict, tags: list[str],
        strategy: str, dedup: Deduplicator, default_status: str, source_label: str,
        file_id: uuid.UUID, counters: dict,
    ) -> str:
        match = await dedup.find_match(session, clean) if strategy != "CREATE_NEW" else None
        is_duplicate = match is not None and match.confidence in DUPLICATE_CONFIDENCES

        if is_duplicate and strategy == "SKIP_DUPLICATES":
            counters["duplicate"] += 1
            return "skipped"

        lead = None
        if is_duplicate and strategy == "UPDATE_EXISTING" and match.lead_id is not None:
            lead = await session.get(Lead, match.lead_id)
            if lead is not None:
                for key, value in clean.items():
                    if key.endswith("_norm") or key == "name_key":
                        continue
                    if value not in (None, "") and (getattr(lead, key) in (None, "")):
                        setattr(lead, key, value)
                for key in ("email_norm", "phone_norm", "website_norm", "name_key"):
                    if clean.get(key):
                        setattr(lead, key, clean[key])
                lead.seen_count = (lead.seen_count or 1) + 1
                lead.last_seen_at = datetime.now(timezone.utc)
                lead.quality_score = compute_quality_score(lead.to_public_dict())
                lead.updated_at = datetime.now(timezone.utc)
                if tags:
                    await self.tags.sync_names(session, lead.id, sorted(
                        set((lead.tags or []) + tags)
                    ), user_id=batch.created_by)
                counters["updated"] += 1
                await self.activities.log(
                    session, lead.id, EVENT_IMPORTED,
                    message=f"Updated by import {batch.filename}",
                    metadata={"import_batch_id": str(batch.id)},
                    user_id=batch.created_by,
                )
                return "updated"

        if is_duplicate and strategy == "REVIEW" and match.lead_id is not None:
            # insert flagged + create the review candidate below (after flush)
            pass

        # create the lead (also the REVIEW path: insert flagged + candidate)
        now = datetime.now(timezone.utc)
        meta = dict(clean.get("metadata") or {})
        if is_duplicate and strategy == "REVIEW" and match is not None:
            meta["possible_duplicate_of"] = str(match.lead_id)
            meta["possible_duplicate_confidence"] = match.confidence.value
        lead = Lead(
            **{k: v for k, v in clean.items() if not k.endswith("_norm") and k != "name_key" and k != "metadata" and k != "tags"},
            email_norm=clean.get("email_norm"),
            phone_norm=clean.get("phone_norm"),
            website_norm=clean.get("website_norm"),
            name_key=clean.get("name_key"),
            source=source_label[:100],
            source_type="import",
            source_url=clean.get("source_url"),
            source_id=clean.get("source_id"),
            imported_file_id=file_id,
            import_batch_id=batch.id,
            status=default_status,
            metadata_json=meta,
            first_seen_at=now,
            last_seen_at=now,
            created_by=batch.created_by,
        )
        lead.quality_score = compute_quality_score(lead.to_public_dict())
        session.add(lead)
        await session.flush()
        if tags:
            await self.tags.sync_names(session, lead.id, tags, user_id=batch.created_by)
        await self.activities.log(
            session, lead.id, EVENT_IMPORTED,
            message=f"Imported from {batch.filename}",
            metadata={"import_batch_id": str(batch.id)},
            user_id=batch.created_by,
        )
        if is_duplicate and strategy == "REVIEW" and match is not None:
            await self.duplicates.record_candidate(
                session, match.lead_id, lead.id,
                confidence=_review_confidence(match), matched_on=match.matched_on,
                origin="import",
            )
            counters["review"] += 1
            return "review"
        if is_duplicate:
            counters["duplicate"] += 1
            return "created_duplicate"
        counters["imported"] += 1
        return "created"

    @staticmethod
    def _row_error(payload: dict) -> str | None:
        has_name = bool(str(payload.get("business_name") or "").strip())
        has_contact = bool(str(payload.get("contact_name") or "").strip())
        if not has_name and not has_contact:
            # a bare email/phone with no name is not a usable business lead
            # (§24 example: "Row 43: Missing required business name")
            return "Missing required business name (and no contact name)"
        return None

    async def _flush(
        self, session: AsyncSession, batch: ImportBatch, counters: dict, error_summary: list[dict]
    ) -> None:
        batch.total_rows = counters["total"]
        batch.valid_rows = counters["valid"]
        batch.invalid_rows = counters["invalid"]
        batch.imported_rows = counters["imported"]
        batch.duplicate_rows = counters["duplicate"]
        batch.updated_rows = counters["updated"]
        batch.review_rows = counters["review"]
        batch.error_summary = error_summary[:1000]
        await session.commit()

    # ---------------------------------------------------------------- worker
    async def claim_next(self, session: AsyncSession, owner: str) -> ImportBatch | None:
        """Atomic-enough DB-first claim: only a row still QUEUED transitions."""
        row = (
            await session.scalars(
                select(ImportBatch)
                .where(ImportBatch.status == ImportStatus.QUEUED.value)
                .order_by(ImportBatch.created_at)
                .limit(1)
            )
        ).first()
        if row is None:
            return None
        result = await session.execute(
            update(ImportBatch)
            .where(ImportBatch.id == row.id, ImportBatch.status == ImportStatus.QUEUED.value)
            .values(status=ImportStatus.PROCESSING.value, lease_owner=owner,
                    leased_at=datetime.now(timezone.utc))
        )
        await session.commit()
        return row if (result.rowcount or 0) else None

    async def cancel(self, session: AsyncSession, batch: ImportBatch) -> ImportBatch:
        if batch.status not in (ImportStatus.QUEUED.value,):
            raise ValidationError(f"Cannot cancel a batch in status {batch.status}")
        batch.status = ImportStatus.CANCELLED.value
        batch.completed_at = datetime.now(timezone.utc)
        await session.commit()
        return batch
